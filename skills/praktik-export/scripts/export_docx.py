#!/usr/bin/env python3
"""Export kurzu PRAKTIK-AI z Markdownu (struktura A0601) do .docx v domácím stylu.

Použití:
    python export_docx.py <kurz.md> [--out <kurz.docx>] [--zdroj "text do zápatí"]

Vyžaduje: pip install python-docx
"""
from __future__ import annotations

import argparse
import re
import sys
from pathlib import Path

try:
    from docx import Document
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.enum.table import WD_TABLE_ALIGNMENT
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Pt, Cm, RGBColor
except ImportError:
    sys.exit("Chybí python-docx: pip install python-docx")

BLUE = RGBColor(0x1F, 0x5C, 0x99)
LIGHT_BLUE_HEX = "DEEAF6"
BORDER_BLUE_HEX = "9DC3E6"
GREY = RGBColor(0x59, 0x59, 0x59)
FONT = "Arial"

BLOCK_H1 = re.compile(r"^#\s+(MODUL|CAPSTONE|ZÁVĚREČNÝ TEST|PROCVIČOVACÍ TEST|ZDROJE|3' VIDEO)", re.I)


def set_cell_bg(cell, hex_color):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hex_color)
    tcPr.append(shd)


def set_par_shading(par, hex_color):
    pPr = par._p.get_or_add_pPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hex_color)
    pPr.append(shd)


def set_left_border(par, hex_color, size=24):
    pPr = par._p.get_or_add_pPr()
    pBdr = OxmlElement("w:pBdr")
    left = OxmlElement("w:left")
    left.set(qn("w:val"), "single")
    left.set(qn("w:sz"), str(size))
    left.set(qn("w:space"), "8")
    left.set(qn("w:color"), hex_color)
    pBdr.append(left)
    pPr.append(pBdr)


def set_bottom_border(par, hex_color, size=6):
    pPr = par._p.get_or_add_pPr()
    pBdr = OxmlElement("w:pBdr")
    b = OxmlElement("w:bottom")
    b.set(qn("w:val"), "single")
    b.set(qn("w:sz"), str(size))
    b.set(qn("w:space"), "1")
    b.set(qn("w:color"), hex_color)
    pBdr.append(b)
    pPr.append(pBdr)


INLINE = re.compile(r"(\*\*[^*]+\*\*|\*[^*]+\*|`[^`]+`)")


def add_inline(par, text, base_italic=False, size=None, color=None, bold=None):
    """Rozloží **tučně**, *kurzívu* a `kód` na runy."""
    for chunk in INLINE.split(text):
        if not chunk:
            continue
        run = par.add_run()
        if chunk.startswith("**") and chunk.endswith("**"):
            run.text = chunk[2:-2]
            run.bold = True
        elif chunk.startswith("*") and chunk.endswith("*") and len(chunk) > 2:
            run.text = chunk[1:-1]
            run.italic = True
        elif chunk.startswith("`") and chunk.endswith("`"):
            run.text = chunk[1:-1]
            run.font.name = "Consolas"
        else:
            run.text = chunk
        if base_italic:
            run.italic = True
        if bold is not None and run.bold is None:
            run.bold = bold
        run.font.name = run.font.name or FONT
        rPr = run._element.get_or_add_rPr()
        rFonts = rPr.find(qn("w:rFonts"))
        if rFonts is None:
            rFonts = OxmlElement("w:rFonts")
            rPr.append(rFonts)
        rFonts.set(qn("w:eastAsia"), run.font.name)
        if size:
            run.font.size = Pt(size)
        if color is not None:
            run.font.color.rgb = color


def is_table_line(line):
    return line.strip().startswith("|") and line.strip().endswith("|")


def parse_table(lines):
    rows = []
    for l in lines:
        l = l.strip()
        if re.match(r"^\|[\s:\-|]+\|$", l):
            continue  # oddělovač
        cells = [c.strip() for c in l.strip("|").split("|")]
        rows.append(cells)
    return rows


def add_table(doc, rows):
    if not rows:
        return
    ncols = max(len(r) for r in rows)
    table = doc.add_table(rows=len(rows), cols=ncols)
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, row in enumerate(rows):
        for j in range(ncols):
            cell = table.cell(i, j)
            cell.text = ""
            par = cell.paragraphs[0]
            txt = row[j] if j < len(row) else ""
            add_inline(par, txt, size=10, bold=True if i == 0 else None)
            if i == 0:
                set_cell_bg(cell, LIGHT_BLUE_HEX)
    doc.add_paragraph()


def add_heading_block(doc, text):
    """Modrý pruh s bílým textem (moduly, capstone, testy, zdroje)."""
    par = doc.add_paragraph()
    par.paragraph_format.space_before = Pt(18)
    par.paragraph_format.space_after = Pt(6)
    set_par_shading(par, "1F5C99")
    run = par.add_run(" " + text.strip())
    run.bold = True
    run.font.size = Pt(14)
    run.font.name = FONT
    run.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
    return par


def add_footer(doc, text):
    section = doc.sections[0]
    footer = section.footer
    par = footer.paragraphs[0]
    par.text = ""
    run = par.add_run(text + "   |   Strana ")
    run.font.size = Pt(8)
    run.font.color.rgb = GREY
    run.font.name = FONT
    # číslo strany
    fld_run = par.add_run()
    fld_run.font.size = Pt(8)
    fld_run.font.color.rgb = GREY
    for tag, txt in (("begin", None), (None, "PAGE"), ("end", None)):
        if tag:
            fc = OxmlElement("w:fldChar")
            fc.set(qn("w:fldCharType"), tag)
            fld_run._r.append(fc)
        else:
            it = OxmlElement("w:instrText")
            it.set(qn("xml:space"), "preserve")
            it.text = txt
            fld_run._r.append(it)


def extract_meta(md):
    kod = re.search(r"Interní kód \(DB\)\s*\|\s*\**([A-Z]\d{4})", md)
    verze = re.search(r"Verze\s+([\d.]+)\s*\|\s*([^|\n]+?)\s*\|", md)
    return (kod.group(1) if kod else "KURZ",
            verze.group(1).strip() if verze else "1.0",
            verze.group(2).strip() if verze else "")


def convert(md_path: Path, out_path: Path, zdroj: str | None):
    md = md_path.read_text(encoding="utf-8")
    kod, verze, datum = extract_meta(md)
    doc = Document()
    # základní styl
    st = doc.styles["Normal"]
    st.font.name = FONT
    st.font.size = Pt(11)
    st.element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
    st.paragraph_format.space_after = Pt(6)
    st.paragraph_format.line_spacing = 1.15
    for s in doc.sections:
        s.left_margin = s.right_margin = s.top_margin = s.bottom_margin = Cm(2.5)

    lines = md.splitlines()
    i = 0
    first_h1 = True
    in_code = False
    while i < len(lines):
        line = lines[i]
        s = line.strip()
        if s.startswith("```"):
            in_code = not in_code
            i += 1
            continue
        if in_code:
            p = doc.add_paragraph()
            r = p.add_run(line)
            r.font.name = "Consolas"
            r.font.size = Pt(9)
            i += 1
            continue
        if not s or s in ("---", "***"):
            i += 1
            continue
        # tabulka
        if is_table_line(s):
            block = []
            while i < len(lines) and is_table_line(lines[i].strip()):
                block.append(lines[i])
                i += 1
            add_table(doc, parse_table(block))
            continue
        # nadpisy
        m = re.match(r"^(#{1,4})\s+(.*)$", s)
        if m:
            level, text = len(m.group(1)), m.group(2).strip()
            if level == 1 and first_h1:
                p = doc.add_paragraph()
                add_inline(p, text, size=20, color=BLUE, bold=True)
                first_h1 = False
            elif level == 1 and BLOCK_H1.match(s):
                add_heading_block(doc, text)
            elif level == 1:
                p = doc.add_paragraph()
                add_inline(p, text, size=16, color=BLUE, bold=True)
            elif level == 2:
                p = doc.add_paragraph()
                p.paragraph_format.space_before = Pt(12)
                add_inline(p, text, size=13, color=BLUE, bold=True)
                if re.match(r"^F[1-4]\b", text):
                    set_bottom_border(p, BORDER_BLUE_HEX)
            else:
                p = doc.add_paragraph()
                add_inline(p, text, size=12, color=GREY, base_italic=True)
            i += 1
            continue
        # blockquote
        if s.startswith(">"):
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Cm(0.8)
            set_par_shading(p, "FFF4CE")
            set_left_border(p, "C9A227")
            add_inline(p, s.lstrip("> ").strip())
            i += 1
            continue
        # seznamy
        m = re.match(r"^(\d+)[.)]\s+(.*)$", s)
        if m:
            p = doc.add_paragraph(style="List Number")
            add_inline(p, m.group(2))
            i += 1
            # podpoložky a) b) c) odsazené
            while i < len(lines) and re.match(r"^\s{2,}[a-d]\)\s+", lines[i]):
                sp = doc.add_paragraph(style="List Bullet 2")
                add_inline(sp, lines[i].strip())
                i += 1
            continue
        m = re.match(r"^[-*•]\s+(.*)$", s)
        if m:
            p = doc.add_paragraph(style="List Bullet")
            add_inline(p, m.group(1))
            i += 1
            continue
        # kurzivní odstavec (celý v *…*)
        if re.fullmatch(r"\*[^*].*[^*]\*", s) and "**" not in s[:2]:
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Cm(0.8)
            p.paragraph_format.space_before = Pt(4)
            set_left_border(p, BORDER_BLUE_HEX)
            add_inline(p, s[1:-1], base_italic=True)
            i += 1
            continue
        # běžný odstavec (sloučit řádky až po prázdný)
        buf = [s]
        i += 1
        while i < len(lines) and lines[i].strip() and not is_table_line(lines[i].strip()) \
                and not re.match(r"^(#{1,4}\s|[-*•]\s|\d+[.)]\s|>|\*[^*].*\*$|```)", lines[i].strip()):
            buf.append(lines[i].strip())
            i += 1
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        add_inline(p, " ".join(buf))

    footer = (f"Zdroj: {zdroj} · " if zdroj else "") + \
        f"Vzniklo v projektu TQ23000092 PRAKTIK-AI (TA ČR, program SIGMA) · {kod} v{verze}, {datum}"
    add_footer(doc, footer)
    doc.save(out_path)
    return kod, verze


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("md")
    ap.add_argument("--out")
    ap.add_argument("--zdroj", default=None, help="Text zdrojového materiálu do zápatí")
    a = ap.parse_args()
    md_path = Path(a.md)
    out = Path(a.out) if a.out else md_path.with_suffix(".docx")
    kod, verze = convert(md_path, out, a.zdroj)
    print(f"OK {out} ({kod} v{verze})")


if __name__ == "__main__":
    main()
